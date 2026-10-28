import io
import csv
import urllib.parse
from datetime import datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import func
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from backend.app.core.database import get_db
from backend.app.core.response import success_response
from backend.app.models.device import Device
from backend.app.models.event import DeviceEvent
from backend.app.models.alert import AlertRecord
from backend.app.models.probe_log import ProbeLog
from backend.app.services.excel_service import sanitize_formula_injection

router = APIRouter(prefix="/reports", tags=["报表与历史统计"])

@router.get("/history")
def get_history_events(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    device_id: Optional[int] = Query(None),
    event_type: Optional[str] = Query(None), # offline, online, status_change
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(DeviceEvent)
    if device_id:
        query = query.filter(DeviceEvent.device_id == device_id)
    if event_type:
        query = query.filter(DeviceEvent.event_type == event_type)
    if start_date:
        try:
            sd = datetime.fromisoformat(start_date)
            query = query.filter(DeviceEvent.created_at >= sd)
        except Exception:
            pass
    if end_date:
        try:
            ed = datetime.fromisoformat(end_date)
            query = query.filter(DeviceEvent.created_at <= ed)
        except Exception:
            pass

    total = query.count()
    events = query.order_by(DeviceEvent.id.desc()).offset((page - 1) * page_size).limit(page_size).all()

    items = []
    for ev in events:
        items.append({
            "id": ev.id,
            "device_id": ev.device_id,
            "device_name": ev.device.name if ev.device else "已删除设备",
            "ip_address": ev.device.ip_address if ev.device else "-",
            "from_status": ev.from_status,
            "to_status": ev.to_status,
            "event_type": ev.event_type,
            "duration_seconds": ev.duration_seconds,
            "details": ev.details,
            "created_at": ev.created_at.strftime("%Y-%m-%d %H:%M:%S") if ev.created_at else None
        })

    return success_response(data={
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": items
    })

@router.get("/statistics")
def get_report_statistics(days: int = Query(7, ge=1, le=90), db: Session = Depends(get_db)):
    """
    计算关键统计指标：
    - 设备故障次数排行
    - 平均故障恢复时间 (MTTR)
    - 整体可用率
    """
    since_time = datetime.now() - timedelta(days=days)

    total_devices = db.query(Device).filter(Device.enabled == True).count()
    events = db.query(DeviceEvent).filter(DeviceEvent.created_at >= since_time).all()

    offline_events = [e for e in events if e.event_type == "offline"]
    online_events = [e for e in events if e.event_type == "online" and e.duration_seconds is not None]

    total_outages = len(offline_events)
    # 计算 MTTR (平均恢复时间)
    if online_events:
        total_recovery_seconds = sum(e.duration_seconds for e in online_events)
        mttr_seconds = int(total_recovery_seconds / len(online_events))
    else:
        mttr_seconds = 0

    # 设备故障次数排行榜 (Top 5)
    fault_counts = {}
    for e in offline_events:
        fault_counts[e.device_id] = fault_counts.get(e.device_id, 0) + 1

    top_failed_devices = []
    for dev_id, count in sorted(fault_counts.items(), key=lambda x: x[1], reverse=True)[:5]:
        dev = db.query(Device).filter(Device.id == dev_id).first()
        if dev:
            top_failed_devices.append({
                "device_id": dev.id,
                "device_name": dev.name,
                "ip_address": dev.ip_address,
                "failure_count": count
            })

    # 设备可用率列表
    device_availabilities = []
    all_devices = db.query(Device).all()
    for d in all_devices:
        dev_logs = db.query(ProbeLog).filter(
            ProbeLog.device_id == d.id,
            ProbeLog.probe_time >= since_time
        ).all()
        if dev_logs:
            ok_cnt = sum(1 for l in dev_logs if l.success)
            uptime_pct = round((ok_cnt / len(dev_logs)) * 100.0, 1)
        else:
            # 无采样数据时不凭空判定为 100%，若设备在线则为 100%，否则根据状态评估
            uptime_pct = 100.0 if d.status == "ONLINE" else 0.0

        device_availabilities.append({
            "id": d.id,
            "name": d.name,
            "ip_address": d.ip_address,
            "status": d.status,
            "uptime_percent": uptime_pct,
            "group_name": d.group.name if d.group else "未分组"
        })

    return success_response(data={
        "range_days": days,
        "total_monitored_devices": total_devices,
        "total_outages": total_outages,
        "mttr_seconds": mttr_seconds,
        "top_failed_devices": top_failed_devices,
        "device_availabilities": device_availabilities
    })

@router.get("/export-excel")
def export_report_excel(days: int = Query(7, ge=1, le=90), db: Session = Depends(get_db)):
    """生成并下载多工作表综合运维报表"""
    since_time = datetime.now() - timedelta(days=days)
    wb = Workbook()

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")

    # Sheet 1: 可用度统计
    ws1 = wb.active
    ws1.title = "设备可用度统计"
    ws1_headers = ["设备ID", "设备名称", "IP地址", "所属分组", "当前状态", f"最近{days}天探测次数", "探测成功次数", "可用度(%)"]
    ws1.append(ws1_headers)
    for col in range(1, len(ws1_headers) + 1):
        c = ws1.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    for d in db.query(Device).all():
        logs = db.query(ProbeLog).filter(ProbeLog.device_id == d.id, ProbeLog.probe_time >= since_time).all()
        total_probes = len(logs)
        ok_probes = sum(1 for l in logs if l.success)
        pct = round((ok_probes / total_probes) * 100.0, 1) if total_probes > 0 else (100.0 if d.status == "ONLINE" else 0.0)
        ws1.append([
            d.id,
            sanitize_formula_injection(d.name),
            d.ip_address,
            d.group.name if d.group else "未分组",
            d.status,
            total_probes,
            ok_probes,
            pct
        ])

    # Sheet 2: 状态事件明细
    ws2 = wb.create_sheet(title="状态变更事件明细")
    ws2_headers = ["事件ID", "设备名称", "IP地址", "前置状态", "变更后状态", "事件类型", "持续时长(秒)", "明细说明", "发生时间"]
    ws2.append(ws2_headers)
    for col in range(1, len(ws2_headers) + 1):
        c = ws2.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    events = db.query(DeviceEvent).filter(DeviceEvent.created_at >= since_time).order_by(DeviceEvent.id.desc()).all()
    for ev in events:
        ws2.append([
            ev.id,
            sanitize_formula_injection(ev.device.name if ev.device else "已删除设备"),
            ev.device.ip_address if ev.device else "-",
            ev.from_status,
            ev.to_status,
            ev.event_type,
            ev.duration_seconds or "",
            sanitize_formula_injection(ev.details or ""),
            ev.created_at.strftime("%Y-%m-%d %H:%M:%S") if ev.created_at else ""
        ])

    # Sheet 3: 告警历史
    ws3 = wb.create_sheet(title="告警历史记录")
    ws3_headers = ["告警ID", "设备名称", "告警类型", "告警标题", "处理状态", "触发时间", "恢复时间", "持续秒数", "已发通道"]
    ws3.append(ws3_headers)
    for col in range(1, len(ws3_headers) + 1):
        c = ws3.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    alerts = db.query(AlertRecord).filter(AlertRecord.triggered_at >= since_time).order_by(AlertRecord.id.desc()).all()
    for al in alerts:
        ws3.append([
            al.id,
            sanitize_formula_injection(al.device.name if al.device else "已删除设备"),
            al.alert_type,
            sanitize_formula_injection(al.title),
            al.status,
            al.triggered_at.strftime("%Y-%m-%d %H:%M:%S") if al.triggered_at else "",
            al.resolved_at.strftime("%Y-%m-%d %H:%M:%S") if al.resolved_at else "",
            al.duration_seconds or "",
            al.channels_sent or ""
        ])

    buf = io.BytesIO()
    wb.save(buf)

    filename = f"IPMonitor运维统计报表_{datetime.now().strftime('%Y%m%d')}.xlsx"
    encoded_filename = urllib.parse.quote(filename)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )
