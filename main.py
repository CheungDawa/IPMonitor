import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import threading
import time
import subprocess
import platform
import smtplib
import ssl
from email.message import EmailMessage
import json
import os
import shutil
import socket
import re
from datetime import datetime, timedelta
import sys
import concurrent.futures
import queue
import urllib.request
import urllib.error
from openpyxl import Workbook, load_workbook
import pystray
from PIL import Image, ImageDraw

# ===================== 全局常量与路径规范 =====================
APP_NAME = "IPMonitor"
APP_VERSION = "1.0.0"
CHECKED_CHAR = '✅ 启用'
UNCHECKED_CHAR = '⏸️ 停用'
REGISTRY_PATH = r"Software\Microsoft\Windows\CurrentVersion\Run"
REG_KEY_NAME = "IPMonitorAutoStart"

# 统一绝对基准目录，避免注册表开机自启时工作目录为 System32 导致的读写异常
BASE_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, 'frozen', False) else __file__))
CONFIG_FILE = os.path.join(BASE_DIR, 'config.json')
CONFIG_EXAMPLE_FILE = os.path.join(BASE_DIR, 'config.example.json')
LOG_FILE = os.path.join(BASE_DIR, "IPMonitor_run.log")

DEFAULT_CONFIG = {
    "monitoring": {
        "interval_seconds": 60,
        "retries_before_offline": 3,
        "timeout_seconds": 3,
        "re_alert_interval_minutes": 60,
        "max_concurrent_workers": 50
    },
    "alert": {
        "desktop_notify": True,
        "webhook_enabled": False,
        "webhook_url": "",
        "webhook_type": "wecom"
    },
    "email": {
        "enabled": False,
        "smtp_server": "smtp.example.com",
        "smtp_port": 465,
        "use_ssl": True,
        "sender_email": "your_email@example.com",
        "password": "your_smtp_auth_token",
        "recipient_email": "recipient@example.com",
        "cc_email": "",
        "quiet_hours_enabled": False,
        "quiet_start_hour": 22,
        "quiet_end_hour": 8
    },
    "system": {"auto_start": False},
    "targets": []
}

# ===================== 通用工具函数 =====================
def get_now_hour(): return datetime.now().hour

def format_duration(duration):
    if not isinstance(duration, timedelta): return ""
    total_sec = int(duration.total_seconds())
    days, rem = divmod(total_sec, 86400)
    hours, rem = divmod(rem, 3600)
    mins, secs = divmod(rem, 60)
    parts = []
    if days > 0: parts.append(f"{days}天")
    if hours > 0: parts.append(f"{hours}小时")
    if mins > 0: parts.append(f"{mins}分")
    return "".join(parts) if parts else f"{secs}秒"

def check_quiet_hour(hour, start, end):
    if start > end: return hour >= start or hour < end
    return start <= hour < end

def set_auto_start(enable: bool):
    if platform.system() != "Windows": return False
    try:
        import winreg
        exe_path = sys.executable if getattr(sys, 'frozen', False) else os.path.abspath(__file__)
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, REGISTRY_PATH, 0, winreg.KEY_WRITE)
        if enable:
            cmd = f'"{exe_path}"' if getattr(sys, 'frozen', False) else f'"{sys.executable}" "{exe_path}"'
            winreg.SetValueEx(key, REG_KEY_NAME, 0, winreg.REG_SZ, cmd)
        else:
            try: winreg.DeleteValue(key, REG_KEY_NAME)
            except FileNotFoundError: pass
        winreg.CloseKey(key)
        return True
    except Exception as e:
        messagebox.showerror("开机自启设置失败", f"注册表操作异常：{str(e)}")
        return False

def get_auto_start_status():
    if platform.system() != "Windows": return False
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, REGISTRY_PATH)
        val = winreg.QueryValueEx(key, REG_KEY_NAME)[0]
        winreg.CloseKey(key)
        return bool(val)
    except Exception: return False

def create_tray_icon():
    img = Image.new("RGB", (64, 64), color=(0, 103, 192))
    draw = ImageDraw.Draw(img)
    draw.ellipse((10, 10, 54, 54), fill=(243, 243, 243), outline=(255, 255, 255), width=2)
    draw.text((16, 20), "IPM", fill=(0, 103, 192), font_size=20)
    return img

def acquire_single_instance_lock():
    if platform.system() == "Windows":
        try:
            import ctypes
            kernel32 = ctypes.windll.kernel32
            mutex = kernel32.CreateMutexW(None, False, f"Global\\{APP_NAME}_SingleInstance_Mutex_2026")
            if kernel32.GetLastError() == 183: # ERROR_ALREADY_EXISTS
                messagebox.showerror("多重实例化阻断", f"进程 {APP_NAME} 已在运行中，请检查系统托盘。")
                sys.exit(0)
            return mutex
        except Exception:
            pass
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.bind(("127.0.0.1", 54321))
        return s
    except OSError:
        messagebox.showerror("多重实例化阻断", f"进程 {APP_NAME} 已在运行中，请检查系统托盘。")
        sys.exit(0)

# ===================== 主程序窗口类 =====================
class IPMonitorApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.state_lock = threading.Lock()
        self.target_states = {}
        self.monitoring_thread = None
        self.monitoring_active = threading.Event()
        self.all_tree_items = {}
        self.tray_icon = None
        self.is_hide_to_tray = False
        
        self.alert_queue = queue.Queue()
        self.alert_worker_thread = threading.Thread(target=self._alert_worker_loop, daemon=True)
        self.alert_worker_thread.start()
        
        self.setup_win11_theme()
        self.app_config = self.load_config()
            
        self.title(f"内网设备离线监控助手 - {APP_VERSION}")
        self.geometry("1150x800")
        self.minsize(960, 650)
        self.configure(bg="#FFFFFF") 
        
        self.create_widgets()
        self.load_targets_to_treeview()
        self.init_tray()
        
        self.tree.bind("<Double-1>", self.on_item_double_click)
        self.tree.bind("<Button-1>", self.on_toggle_enabled)
        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.app_config["system"]["auto_start"] = get_auto_start_status()

    def setup_win11_theme(self):
        style = ttk.Style(self)
        if 'clam' in style.theme_names():
            style.theme_use('clam')
            
        win_bg = "#FFFFFF"
        text_color = "#333333"
        sub_text_color = "#666666"
        border_color = "#E0E0E0"
        primary_color = "#2563EB"
        
        font_family = "Microsoft YaHei UI" if platform.system() == "Windows" else "Microsoft YaHei"
        global_font = (font_family, 10)
        
        style.configure(".", font=global_font, background=win_bg, foreground=text_color)
        
        # Primary Button
        style.configure("Primary.TButton", font=(font_family, 10, "bold"), background=primary_color, foreground="white", bordercolor=primary_color, lightcolor=primary_color, darkcolor=primary_color, padding=(12, 6))
        style.map("Primary.TButton", background=[('active', '#1D4ED8'), ('pressed', '#1E40AF')])
        
        # Standard Button
        style.configure("TButton", font=global_font, background="#F8F9FA", foreground=text_color, bordercolor=border_color, focuscolor="", padding=(10, 5))
        style.map("TButton", background=[('active', '#F1F3F5'), ('pressed', '#E9ECEF')], bordercolor=[('active', '#D1D5DB')])
        
        # Entry
        style.configure("TEntry", padding=6, fieldbackground="#F9FAFB", bordercolor=border_color, selectbackground=primary_color)
        style.map("TEntry", bordercolor=[('focus', primary_color)])
        
        # Treeview
        style.configure("Treeview", font=(font_family, 10), rowheight=45, background=win_bg, fieldbackground=win_bg, borderwidth=0)
        style.configure("Treeview.Heading", font=(font_family, 10, "bold"), background="#F9FAFB", foreground=sub_text_color, padding=8, borderwidth=0)
        style.map('Treeview', background=[('selected', '#EFF6FF')], foreground=[('selected', text_color)])
        style.layout("Treeview", [('Treeview.treearea', {'sticky': 'nswe'})])

    def deep_fix_config(self, target: dict, default: dict):
        for k, def_val in default.items():
            if k not in target:
                target[k] = def_val
                continue
            curr_val = target[k]
            if isinstance(curr_val, list):
                continue
            if isinstance(def_val, dict) and isinstance(curr_val, dict):
                self.deep_fix_config(curr_val, def_val)

    def load_config(self):
        if not os.path.exists(CONFIG_FILE) and os.path.exists(CONFIG_EXAMPLE_FILE):
            try:
                shutil.copyfile(CONFIG_EXAMPLE_FILE, CONFIG_FILE)
            except Exception: pass

        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, 'r', encoding='utf-8') as f: cfg = json.load(f)
                if not isinstance(cfg, dict): raise ValueError()
                self.deep_fix_config(cfg, DEFAULT_CONFIG)
                return cfg
            except Exception: pass
        return DEFAULT_CONFIG

    def save_config(self):
        with self.state_lock:
            self.app_config['targets'] = [v['data'] for v in self.all_tree_items.values()]
        try:
            with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
                json.dump(self.app_config, f, indent=4, ensure_ascii=False)
            self.log("配置信息已同步保存至本地。")
        except Exception as e:
            self.log(f"配置保存失败：{str(e)}")

    def log(self, msg):
        full_time = time.strftime('%Y-%m-%d %H:%M:%S')
        try:
            with open(LOG_FILE, "a", encoding="utf-8") as f:
                f.write(f"[{full_time}] {msg}\n")
        except Exception:
            pass

        def _log():
            try:
                self.log_text.config(state=tk.NORMAL)
                self.log_text.insert(tk.END, f"[{time.strftime('%H:%M:%S')}] {msg}\n")
                lines = int(self.log_text.index('end-1c').split('.')[0])
                if lines > 500:
                    self.log_text.delete('1.0', f'{lines - 500 + 1}.0')
                self.log_text.see(tk.END)
                self.log_text.config(state=tk.DISABLED)
            except Exception: pass
            
        self.after(0, _log)

    def init_tray(self):
        menu = pystray.Menu(
            pystray.MenuItem("打开主面板", self.tray_show_window, default=True), 
            pystray.MenuItem("退出应用", self.tray_quit_app)
        )
        self.tray_icon = pystray.Icon("ipmonitor", create_tray_icon(), APP_NAME, menu)
        threading.Thread(target=self.tray_icon.run, daemon=True).start()

    def tray_show_window(self):
        self.is_hide_to_tray = False
        self.deiconify()
        self.lift()

    def tray_quit_app(self): self.on_closing()

    def on_closing(self):
        res = messagebox.askyesnocancel("退出程序", "是否将程序最小化到系统托盘后台运行？\n\n[是] 隐藏到托盘持续监控\n[否] 完全退出并停止监控\n[取消] 撤销操作", parent=self)
        if res is True:
            self.hide_to_tray()
        elif res is False:
            self.stop_monitoring()
            if self.tray_icon: self.tray_icon.stop()
            self.save_config()
            self.destroy()

    def hide_to_tray(self):
        self.is_hide_to_tray = True
        self.withdraw()
        self.log("程序已挂起至后台托盘，持续监控中...")

    def create_widgets(self):
        # Sidebar
        self.sidebar = tk.Frame(self, bg="#F3F5F7")
        self.sidebar.pack(side=tk.LEFT, fill=tk.Y)
        
        # Logo area
        logo_frm = tk.Frame(self.sidebar, bg="#F3F5F7", pady=25)
        logo_frm.pack(fill=tk.X)
        tk.Label(logo_frm, text="🛡️", font=("Segoe UI Emoji", 24), bg="#F3F5F7", fg="#2563EB").pack(side=tk.LEFT, padx=(20, 5))
        tk.Label(logo_frm, text="IP Monitor", font=("Microsoft YaHei UI", 13, "bold"), bg="#F3F5F7", fg="#1F2937").pack(side=tk.LEFT, padx=(0, 25))
        
        # Content Area
        self.content_area = tk.Frame(self, bg="#FFFFFF")
        self.content_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        # ----------------- Dashboard Tab -----------------
        self.tab_dashboard = tk.Frame(self.content_area, bg="#FFFFFF", padx=30, pady=25)
        
        dash_top = tk.Frame(self.tab_dashboard, bg="#FFFFFF")
        dash_top.pack(fill=tk.X, pady=(0, 20))
        tk.Label(dash_top, text="节点监控", font=("Microsoft YaHei UI", 18, "bold"), bg="#FFFFFF", fg="#111827").pack(side=tk.LEFT)
        
        search_frm = tk.Frame(dash_top, bg="#FFFFFF")
        search_frm.pack(side=tk.RIGHT)
        tk.Label(search_frm, text="🔍", bg="#FFFFFF", fg="#6B7280").pack(side=tk.LEFT, padx=(0, 5))
        self.filter_var = tk.StringVar()
        filter_entry = ttk.Entry(search_frm, textvariable=self.filter_var, width=30)
        filter_entry.pack(side=tk.LEFT)
        self.filter_var.trace_add("write", self.filter_list)
        
        action_bar = tk.Frame(self.tab_dashboard, bg="#FFFFFF")
        action_bar.pack(fill=tk.X, pady=(0, 15))
        
        self.start_button = ttk.Button(action_bar, text="▶ 启动监控", command=self.start_monitoring, style="Primary.TButton")
        self.start_button.pack(side=tk.LEFT, padx=(0, 10))
        self.stop_button = ttk.Button(action_bar, text="■ 停止", command=self.stop_monitoring, state=tk.DISABLED)
        self.stop_button.pack(side=tk.LEFT, padx=(0, 20))
        
        self.add_button = ttk.Button(action_bar, text="➕ 添加", command=self.add_target)
        self.add_button.pack(side=tk.LEFT, padx=(0, 8))
        self.delete_button = ttk.Button(action_bar, text="➖ 删除", command=self.delete_selected_targets)
        self.delete_button.pack(side=tk.LEFT, padx=(0, 20))
        
        ttk.Button(action_bar, text="⬇️ 导出", command=self.export_excel_template).pack(side=tk.RIGHT, padx=(10, 0))
        ttk.Button(action_bar, text="⬆️ 导入", command=self.import_from_excel).pack(side=tk.RIGHT)
        
        tree_card = tk.Frame(self.tab_dashboard, bg="#FFFFFF", highlightbackground="#E5E7EB", highlightthickness=1)
        tree_card.pack(fill=tk.BOTH, expand=True)
        
        self.tree = ttk.Treeview(tree_card, columns=("seq", "enabled", "name", "address", "method", "port", "status", "note"), show="headings")
        col_def = {
            "seq": ("#", 50, "center"),
            "enabled": ("状态", 80, "center"),
            "name": ("资源名称", 180, "w"),
            "address": ("地址", 160, "w"),
            "method": ("协议", 80, "center"),
            "port": ("端口", 80, "center"),
            "status": ("链路状态", 140, "center"),
            "note": ("备注", 200, "w")
        }
        for cid, (text, w, anchor) in col_def.items():
            self.tree.heading(cid, text=text, command=lambda _c=cid: self.sort_list(_c, False))
            self.tree.column(cid, width=w, anchor=anchor, stretch=False if cid in ("seq", "enabled", "method", "port") else True)
            
        self.tree.tag_configure('online', foreground='#059669', background='#ECFDF5') 
        self.tree.tag_configure('offline', foreground='#DC2626', background='#FEF2F2') 
        self.tree.tag_configure('unknown', foreground='#6B7280') 
        self.tree.tag_configure('checking', foreground='#2563EB', background='#EFF6FF')
        self.tree.tag_configure('disabled', foreground='#9CA3AF')
        
        self.tree.pack(fill=tk.BOTH, expand=True, side=tk.LEFT)
        scroll = ttk.Scrollbar(tree_card, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        
        # ----------------- Logs Tab -----------------
        self.tab_logs = tk.Frame(self.content_area, bg="#FFFFFF", padx=30, pady=25)
        
        log_top = tk.Frame(self.tab_logs, bg="#FFFFFF")
        log_top.pack(fill=tk.X, pady=(0, 20))
        tk.Label(log_top, text="运行日志", font=("Microsoft YaHei UI", 18, "bold"), bg="#FFFFFF", fg="#111827").pack(side=tk.LEFT)
        
        ttk.Button(log_top, text="🗑️ 清空日志", command=lambda: [self.log_text.config(state=tk.NORMAL), self.log_text.delete('1.0', tk.END), self.log_text.config(state=tk.DISABLED)]).pack(side=tk.RIGHT)
        
        log_card = tk.Frame(self.tab_logs, bg="#FFFFFF", highlightbackground="#E5E7EB", highlightthickness=1)
        log_card.pack(fill=tk.BOTH, expand=True)
        
        self.log_text = scrolledtext.ScrolledText(log_card, state=tk.DISABLED, bg="#F9FAFB", fg="#374151", font=("Consolas", 10), borderwidth=0, padx=10, pady=10)
        self.log_text.pack(fill=tk.BOTH, expand=True)

        # Nav Variables & Tab Switch
        self.nav_var = tk.StringVar(value="dashboard")
        
        def switch_tab():
            tab = self.nav_var.get()
            if tab == "dashboard":
                self.tab_logs.pack_forget()
                self.tab_dashboard.pack(fill=tk.BOTH, expand=True)
            elif tab == "logs":
                self.tab_dashboard.pack_forget()
                self.tab_logs.pack(fill=tk.BOTH, expand=True)
                
        def make_nav_btn(val, text, icon):
            rb = tk.Radiobutton(self.sidebar, text=f"{icon}   {text}", variable=self.nav_var, value=val, command=switch_tab,
                                indicatoron=0, bg="#F3F5F7", activebackground="#E5E7EB", selectcolor="#E5E7EB", 
                                fg="#4B5563", font=("Microsoft YaHei UI", 11), borderwidth=0, relief="flat", anchor="w", padx=25, pady=12, cursor="hand2")
            rb.pack(fill=tk.X, pady=2, padx=10)
            
        make_nav_btn("dashboard", "仪表盘", "📊")
        make_nav_btn("logs", "运行日志", "📝")
        
        # Bottom sidebar buttons
        bottom_frm = tk.Frame(self.sidebar, bg="#F3F5F7")
        bottom_frm.pack(side=tk.BOTTOM, fill=tk.X, pady=25, padx=10)
        
        tk.Button(bottom_frm, text="⚙️   设置", command=self.open_settings_window, bg="#F3F5F7", activebackground="#E5E7EB", fg="#4B5563", font=("Microsoft YaHei UI", 11), borderwidth=0, relief="flat", anchor="w", padx=25, pady=10, cursor="hand2").pack(fill=tk.X)
        tk.Button(bottom_frm, text="ℹ️   关于", command=self.show_about, bg="#F3F5F7", activebackground="#E5E7EB", fg="#4B5563", font=("Microsoft YaHei UI", 11), borderwidth=0, relief="flat", anchor="w", padx=25, pady=10, cursor="hand2").pack(fill=tk.X)
        
        switch_tab()

    def show_about(self):
        about_text = (
            f"软件名称：{APP_NAME} (内网设备离线监控助手)\n"
            f"软件版本：{APP_VERSION}\n"
            f"开源协议：MIT License\n\n"
            "【开发者信息】\n"
            "开发者：张定荣\n"
            "联系邮箱：DawaCheung@163.com\n\n"
            "【项目介绍】\n"
            "轻量级、跨平台的内网设备存活探测与离线告警小工具，\n"
            "支持 ICMP Ping、TCP 端口及本地 ARP 快速探活，\n"
            "支持系统托盘桌面气泡、企业微信/钉钉/飞书 Webhook 及 SMTP 邮件告警。"
        )
        messagebox.showinfo("关于 IPMonitor", about_text, parent=self)

    def on_toggle_enabled(self, event):
        region = self.tree.identify_region(event.x, event.y)
        if region != "cell": return
        item_id = self.tree.identify_row(event.y)
        col_id = self.tree.identify_column(event.x)
        if col_id == '#2' and item_id in self.all_tree_items:
            with self.state_lock:
                data = self.all_tree_items[item_id]['data']
                data['enabled'] = not data['enabled']
                self.tree.set(item_id, 'enabled', CHECKED_CHAR if data['enabled'] else UNCHECKED_CHAR)
                if not data['enabled']:
                    self.target_states[item_id]['status'] = "已挂起"
                    self.update_target_status_ui(item_id, "⏸️ 已挂起", ('disabled',))
                else:
                    self.target_states[item_id]['status'] = "未知"
                    self.update_target_status_ui(item_id, "⚪ 待初始化", ('unknown',))
            self.save_config()

    def load_targets_to_treeview(self):
        with self.state_lock:
            self.tree.delete(*self.tree.get_children())
            self.target_states.clear()
            self.all_tree_items.clear()
            for idx, target in enumerate(self.app_config.get('targets', [])):
                iid = f"I{idx+1:04}"
                enable = target.get('enabled', True)
                status_txt = "未知" if enable else "已挂起"
                ui_status = "⚪ 未知" if enable else "⏸️ 已挂起"
                vals = (idx + 1, CHECKED_CHAR if enable else UNCHECKED_CHAR, target['name'], target['address'], target['method'], target.get('port', ''), ui_status, target.get('note', ''))
                tag = 'unknown' if enable else 'disabled'
                self.tree.insert("", tk.END, iid=iid, values=vals, tags=(tag,))
                self.target_states[iid] = {"status": status_txt, "failures": 0, "offline_since": None, "last_alert_time": None, "ui_status": ui_status}
                self.all_tree_items[iid] = {'data': target}

    def filter_list(self, *args):
        q = self.filter_var.get().lower()
        with self.state_lock:
            self.tree.delete(*self.tree.get_children())
            for idx, (iid, item) in enumerate(self.all_tree_items.items()):
                d = item['data']
                if q in d['name'].lower() or q in d['address'].lower() or q in d.get('note','').lower():
                    en = d.get('enabled', True)
                    real_st = self.target_states[iid]['status']
                    vals = (idx + 1, CHECKED_CHAR if en else UNCHECKED_CHAR, d['name'], d['address'], d['method'], d.get('port',''), self.target_states[iid]['ui_status'], d.get('note',''))
                    tag = {"在线":"online","离线":"offline","已挂起":"disabled","未知":"unknown","检测中":"checking"}.get(real_st, 'unknown')
                    self.tree.insert("", tk.END, iid=iid, values=vals, tags=(tag,))

    def sort_list(self, col, reverse):
        items = [(self.tree.set(k, col), k) for k in self.tree.get_children('')]
        items.sort(reverse=reverse)
        for idx, (_, k) in enumerate(items): self.tree.move(k, '', idx)
        self.tree.heading(col, command=lambda: self.sort_list(col, not reverse))

    def update_target_status_ui(self, item_id, status_text, tags):
        with self.state_lock:
            if item_id in self.target_states:
                self.target_states[item_id]['ui_status'] = status_text
                
        def _update():
            if self.tree.exists(item_id):
                try:
                    vals = list(self.tree.item(item_id, 'values'))
                    vals[6] = status_text
                    self.tree.item(item_id, values=tuple(vals), tags=tags)
                except Exception: pass
        self.after(0, _update)

    def add_target(self):
        dialog = AddEditWindow(self, "添加监测节点")
        if dialog.result:
            with self.state_lock:
                existing_ids = [int(k[1:]) for k in self.all_tree_items.keys()]
                next_idx = max(existing_ids) + 1 if existing_ids else 1
                iid = f"I{next_idx:04}"
                self.all_tree_items[iid] = {'data': dialog.result}
                en = dialog.result.get('enabled', True)
                ui_status = "⚪ 未知" if en else "⏸️ 已挂起"
                self.target_states[iid] = {"status": "未知" if en else "已挂起", "failures":0, "offline_since":None, "last_alert_time":None, "ui_status": ui_status}
            self.filter_list()
            self.log(f"资源录入成功：{dialog.result['name']}")
            self.save_config()

    def on_item_double_click(self, event):
        if self.tree.identify_region(event.x, event.y) != "cell" or self.tree.identify_column(event.x) in ('#1', '#2'): return
        iid = self.tree.focus()
        if not iid: return
        with self.state_lock:
            init_data = self.all_tree_items[iid]['data']
        dialog = AddEditWindow(self, "编辑监测节点", init_data)
        if dialog.result:
            with self.state_lock:
                self.all_tree_items[iid]['data'] = dialog.result
                en = dialog.result.get('enabled', True)
                self.target_states[iid]['status'] = "未知" if en else "已挂起"
                self.target_states[iid]['ui_status'] = "⚪ 未知" if en else "⏸️ 已挂起"
            self.filter_list()
            self.log(f"配置修改下发：{dialog.result['name']}")
            self.save_config()

    def delete_selected_targets(self):
        sel = self.tree.selection()
        if not sel: return messagebox.showwarning("操作提示", "请在列表中选中需要移除的数据行。")
        if not messagebox.askyesno("移除确认", f"系统将永久移除选定的 {len(sel)} 个节点，是否继续？"): return
        with self.state_lock:
            for iid in sel:
                if iid in self.all_tree_items:
                    name = self.all_tree_items[iid]['data']['name']
                    del self.all_tree_items[iid]
                    if iid in self.target_states:
                        del self.target_states[iid]
                    self.log(f"已清理节点：{name}")
        self.filter_list()
        self.save_config()

    def export_excel_template(self):
        path = filedialog.asksaveasfilename(title="保存台账模板", defaultextension=".xlsx", filetypes=[("Excel 文件","*.xlsx")], initialfile="Monitor_Assets.xlsx")
        if not path: return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Assets"
            ws.append(["Enabled", "Name", "Address", "Method", "Port", "Note"])
            ws.append([True, "业务网关", "192.168.1.1", "ping", "", "核心机房出口"])
            ws.append([False, "测试服务器", "192.168.1.2", "tcp", "445", "SMB 共享"])
            wb.save(path)
            self.log(f"导出向导完成：{path}")
            messagebox.showinfo("导出成功", f"模板已生成至：\n{path}")
        except Exception as e:
            messagebox.showerror("导出失败", f"生成 Excel 文件出错：{str(e)}")

    def import_from_excel(self):
        path = filedialog.askopenfilename(title="读取 Excel 台账", filetypes=[("Excel 文件","*.xlsx")])
        if not path: return
        try:
            wb = load_workbook(path, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return messagebox.showwarning("格式提示", "Excel 表格为空或无有效数据行。")
            
            header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
            req_cols = {"name", "address", "method"}
            if not req_cols.issubset(set(header)):
                return messagebox.showerror("格式异常", f"表格缺失必填列：{', '.join(req_cols)}")
            
            if not messagebox.askyesno("覆盖警告", "批量导入将重置当前监控列表，此操作不可逆。是否继续？"):
                return
            
            col_map = {name: idx for idx, name in enumerate(header)}
            new_targets = []
            for r in rows[1:]:
                if not any(r): continue
                name_idx = col_map.get("name")
                addr_idx = col_map.get("address")
                method_idx = col_map.get("method")
                port_idx = col_map.get("port")
                note_idx = col_map.get("note")
                en_idx = col_map.get("enabled")
                
                name_val = str(r[name_idx]).strip() if (name_idx is not None and r[name_idx] is not None) else ""
                addr_val = str(r[addr_idx]).strip() if (addr_idx is not None and r[addr_idx] is not None) else ""
                if not name_val or not addr_val: continue
                
                method_val = str(r[method_idx]).strip().lower() if (method_idx is not None and r[method_idx] is not None) else "ping"
                if method_val not in ("ping", "tcp", "arp"): method_val = "ping"
                
                port_val = ""
                if port_idx is not None and r[port_idx] is not None:
                    try:
                        port_val = str(int(float(r[port_idx])))
                    except Exception:
                        port_val = str(r[port_idx]).strip()
                        
                note_val = str(r[note_idx]).strip() if (note_idx is not None and r[note_idx] is not None) else ""
                
                en_val = True
                if en_idx is not None and r[en_idx] is not None:
                    raw_en = str(r[en_idx]).strip().lower()
                    en_val = raw_en in ("true", "1", "yes", "是", "启用")
                    
                new_targets.append({
                    "name": name_val,
                    "address": addr_val,
                    "method": method_val,
                    "port": port_val if method_val == "tcp" else "",
                    "note": note_val,
                    "enabled": en_val
                })
                
            self.app_config['targets'] = new_targets
            self.load_targets_to_treeview()
            self.save_config()
            self.log(f"数据加载完成，共注入 {len(new_targets)} 条资产记录。")
        except Exception as e:
            messagebox.showerror("读取崩溃", f"导入文件失败：{str(e)}")

    def start_monitoring(self):
        if not self.target_states: return messagebox.showwarning("策略限制", "监控池为空，请先录入目标资产。")
        self.monitoring_active.set()
        self.monitoring_thread = threading.Thread(target=self._monitoring_loop, daemon=True)
        self.monitoring_thread.start()
        self.start_button.config(state=tk.DISABLED)
        self.stop_button.config(state=tk.NORMAL)
        self.tree.unbind("<Double-1>")
        self.tree.unbind("<Button-1>")
        self.log(f"▶ 异步检测队列就绪 (线程池上限: {self.app_config['monitoring'].get('max_concurrent_workers', 50)})")

    def stop_monitoring(self):
        self.monitoring_active.clear()
        if self.monitoring_thread and self.monitoring_thread.is_alive():
            self.monitoring_thread.join(timeout=1.5)
        self.start_button.config(state=tk.NORMAL)
        self.stop_button.config(state=tk.DISABLED)
        self.tree.bind("<Double-1>", self.on_item_double_click)
        self.tree.bind("<Button-1>", self.on_toggle_enabled)
        self.log("■ 调度引擎已安全终止。")

    def _monitoring_loop(self):
        mon_cfg = self.app_config['monitoring']
        interval = mon_cfg.get('interval_seconds', 60)
        retry_cnt = mon_cfg.get('retries_before_offline', 3)
        re_alert_gap = timedelta(minutes=mon_cfg.get('re_alert_interval_minutes', 60))
        max_workers = mon_cfg.get('max_concurrent_workers', 50)

        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            while self.monitoring_active.is_set():
                with self.state_lock:
                    item_list = list(self.target_states.keys())
                
                futures = [executor.submit(self._check_single_target, item_id, retry_cnt, re_alert_gap) 
                           for item_id in item_list if self.monitoring_active.is_set()]
                
                concurrent.futures.wait(futures)
                
                for _ in range(interval):
                    if not self.monitoring_active.is_set(): break
                    time.sleep(1)

    def _check_single_target(self, item_id, retry_cnt, re_alert_gap):
        if not self.monitoring_active.is_set(): return
        with self.state_lock:
            if item_id not in self.all_tree_items or item_id not in self.target_states: return
            target = dict(self.all_tree_items[item_id]['data'])
        if not target.get('enabled', True): return
        
        self.update_target_status_ui(item_id, "🔵 验证中...", ('checking',))
        online = self.check_target(target)
        if not self.monitoring_active.is_set(): return

        with self.state_lock:
            if item_id not in self.target_states: return
            state = self.target_states[item_id]
            if online:
                state['failures'] = 0
                if state['status'] == "离线":
                    state['status'] = "在线"
                    offline_dur = datetime.now() - state['offline_since'] if state['offline_since'] else None
                    self.log(f"🟢 恢复：[{target['name']}] 已重新上线")
                    self.send_alert(target, "恢复在线", offline_dur)
                    state['offline_since'] = None
                    state['last_alert_time'] = None
                self.update_target_status_ui(item_id, "🟢 在线", ('online',))
            else:
                state['failures'] += 1
                if state['failures'] >= retry_cnt:
                    if state['status'] != "离线":
                        state['status'] = "离线"
                        state['offline_since'] = datetime.now()
                        state['last_alert_time'] = datetime.now()
                        self.update_target_status_ui(item_id, "🔴 离线", ('offline',))
                        self.log(f"🔴 告警：[{target['name']}] 确认离线")
                        self.send_alert(target, "确认离线", timedelta(seconds=0))
                    else:
                        if state['last_alert_time'] and datetime.now() - state['last_alert_time'] > re_alert_gap:
                            state['last_alert_time'] = datetime.now()
                            dur = datetime.now() - state['offline_since']
                            self.log(f"🔔 重复告警：[{target['name']}] 持续离线 {format_duration(dur)}")
                            self.send_alert(target, "持续离线", dur)
                    dur = datetime.now() - state['offline_since'] if state['offline_since'] else timedelta(0)
                    self.update_target_status_ui(item_id, f"🔴 离线 ({format_duration(dur)})", ('offline',))
                else:
                    self.update_target_status_ui(item_id, f"🟡 丢包重试 ({state['failures']}/{retry_cnt})", ('checking',))

    def check_target(self, target):
        method = target.get('method', 'ping').lower()
        if method == "ping": return self.ping_check(target['address'])
        elif method == "tcp":
            try: return self.tcp_check(target['address'], int(target.get('port', 0)))
            except Exception: return False
        elif method == "arp": return self.arp_check(target['address'])
        return False

    def ping_check(self, addr):
        timeout_sec = self.app_config['monitoring'].get('timeout_seconds', 3)
        timeout_ms = timeout_sec * 1000
        win = platform.system().lower() == "windows"
        cmd = ["ping", "-n", "1", "-w", str(timeout_ms), addr] if win else ["ping", "-c", "1", "-W", str(timeout_sec), addr]
        si = None
        if win:
            si = subprocess.STARTUPINFO()
            si.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            si.wShowWindow = subprocess.SW_HIDE
        try:
            res = subprocess.run(
                cmd, startupinfo=si, capture_output=True, text=True,
                encoding="gbk", errors="ignore", timeout=timeout_sec + 2
            )
            if win:
                out = res.stdout.lower()
                return (res.returncode == 0 and "ttl=" in out 
                        and "unreachable" not in out and "无法访问" not in out 
                        and "超时" not in out and "timed out" not in out)
            else:
                return res.returncode == 0
        except Exception:
            return False

    def arp_check(self, addr):
        if platform.system() != "Windows": return False
        # 1. 优先调用 Windows 原生 iphlpapi.SendARP (精准快速，杜绝子进程与过期ARP缓存误判)
        try:
            import ctypes
            dest_ip = ctypes.c_ulong(int.from_bytes(socket.inet_aton(addr), 'little'))
            src_ip = ctypes.c_ulong(0)
            mac = (ctypes.c_ubyte * 6)()
            mac_len = ctypes.c_ulong(6)
            res = ctypes.windll.iphlpapi.SendARP(dest_ip, src_ip, ctypes.byref(mac), ctypes.byref(mac_len))
            return res == 0
        except Exception:
            pass
        # 2. 兜底回退 ping+arp 子进程
        try:
            si = subprocess.STARTUPINFO()
            si.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            si.wShowWindow = subprocess.SW_HIDE
            subprocess.run(["ping", "-n", "1", "-w", "500", addr], capture_output=True, startupinfo=si, stderr=subprocess.DEVNULL)
            res = subprocess.run(["arp", "-a", addr], capture_output=True, text=True, encoding="gbk", errors="ignore", stderr=subprocess.DEVNULL, timeout=3)
            return re.search(r"([0-9a-f]{2}-){5}[0-9a-f]{2}", res.stdout.lower()) is not None
        except Exception:
            return False

    def tcp_check(self, addr, port):
        try:
            timeout_sec = self.app_config['monitoring'].get('timeout_seconds', 3)
            with socket.create_connection((addr, int(port)), timeout=timeout_sec):
                return True
        except Exception:
            return False

    def send_alert(self, target, event_type, offline_dur=None):
        name, addr, note = target['name'], target['address'], target.get('note', '')
        dur_str = format_duration(offline_dur)
        now = time.strftime("%Y-%m-%d %H:%M:%S")
        note_text = f"\n备注：{note}" if note else ""
        if event_type == "确认离线":
            subj = f"🔴 [告警] {name} 设备离线"
            body = f"告警类型：设备离线\n设备名称：{name}\nIP地址：{addr}{note_text}\n告警时间：{now}"
        elif event_type == "持续离线":
            subj = f"🔴 [持续告警] {name} 离线{dur_str}"
            body = f"告警类型：持续离线\n设备名称：{name}\nIP地址：{addr}\n已离线时长：{dur_str}{note_text}\n更新时间：{now}"
        elif event_type == "恢复在线":
            subj = f"🟢 [恢复通知] {name} 已重新上线"
            body = f"通知类型：设备恢复\n设备名称：{name}\nIP地址：{addr}{note_text}\n总计离线时长：{dur_str}\n恢复时间：{now}"
        else:
            return
            
        if target.get('method') == "tcp" and "离线" in event_type:
            body += f"\n检测端口：{target.get('port', '')}"

        # 1. 桌面气泡通知
        alert_cfg = self.app_config.get("alert", {})
        if alert_cfg.get("desktop_notify", True) and self.tray_icon:
            try:
                self.tray_icon.notify(f"{name} ({addr}) - {event_type}", title=subj)
            except Exception:
                pass

        # 2. 压入推送队列 (处理 Webhook 与 邮件)
        self.alert_queue.put((subj, body, event_type, target))

    def _alert_worker_loop(self):
        while True:
            try:
                item = self.alert_queue.get()
                alerts = [item]
                
                # 批量汇聚瞬间涌入的告警 (最多合并 20 条，防止大面积掉线引发的告警风暴)
                time.sleep(0.5)
                while len(alerts) < 20 and not self.alert_queue.empty():
                    alerts.append(self.alert_queue.get_nowait())
                    
                alert_cfg = self.app_config.get("alert", {})
                email_cfg = self.app_config.get("email", {})
                
                if len(alerts) == 1:
                    subj, body, _, _ = alerts[0]
                else:
                    subj = f"🔴 [合并告警] 检测到 {len(alerts)} 个设备状态变更"
                    body = f"以下是集中状态变更通知 (共 {len(alerts)} 条)：\n" + "="*40 + "\n\n"
                    for s, b, _, _ in alerts:
                        body += f"【{s}】\n{b}\n" + "-"*40 + "\n"
                        
                # 调度 Webhook 群机器人通知
                if alert_cfg.get("webhook_enabled", False):
                    self._send_webhook(subj, body)
                    
                # 调度邮件通知
                if email_cfg.get("enabled", False):
                    if email_cfg.get('quiet_hours_enabled', False) and check_quiet_hour(get_now_hour(), email_cfg.get('quiet_start_hour', 22), email_cfg.get('quiet_end_hour', 8)):
                        self.log("邮件告警：当前处于免打扰时段，跳过发送")
                    else:
                        self._send_email(subj, body)
                        
                time.sleep(2)
            except Exception as e:
                self.log(f"告警队列调度异常：{str(e)}")

    def _send_webhook(self, title, content):
        alert_cfg = self.app_config.get("alert", {})
        url = alert_cfg.get("webhook_url", "").strip()
        if not url or not url.startswith("http"):
            return
        wtype = alert_cfg.get("webhook_type", "wecom")
        
        full_text = f"{title}\n\n{content}"
        if wtype == "feishu":
            payload = {"msg_type": "text", "content": {"text": full_text}}
        else: # wecom or dingtalk
            payload = {"msgtype": "text", "text": {"content": full_text}}
            
        try:
            req_data = json.dumps(payload).encode("utf-8")
            req = urllib.request.Request(
                url, data=req_data,
                headers={"Content-Type": "application/json; charset=utf-8"},
                method="POST"
            )
            with urllib.request.urlopen(req, timeout=8) as resp:
                if resp.getcode() == 200:
                    self.log(f"Webhook [{wtype}] 告警推送成功：{title}")
                else:
                    self.log(f"Webhook 推送响应码：HTTP {resp.getcode()}")
        except Exception as e:
            self.log(f"Webhook 推送失败：{str(e)}")

    def _send_email(self, subj, body):
        cfg = self.app_config.get('email', {})
        if "example.com" in cfg.get('sender_email', '') or not cfg.get('sender_email'): return
        msg = EmailMessage()
        msg.set_content(body)
        msg["Subject"], msg["From"], msg["To"] = subj, cfg['sender_email'], cfg['recipient_email']
        if cfg.get('cc_email', '').strip():
            msg["Cc"] = ", ".join([x.strip() for x in cfg['cc_email'].replace(";", ",").split(",") if x.strip()])
        for attempt in range(3):
            try:
                ctx = ssl.create_default_context()
                if cfg.get('use_ssl', True):
                    with smtplib.SMTP_SSL(cfg['smtp_server'], int(cfg['smtp_port']), context=ctx, timeout=10) as srv:
                        srv.login(cfg['sender_email'], cfg['password'])
                        srv.send_message(msg)
                else:
                    with smtplib.SMTP(cfg['smtp_server'], int(cfg['smtp_port']), timeout=10) as srv:
                        srv.starttls(context=ctx)
                        srv.login(cfg['sender_email'], cfg['password'])
                        srv.send_message(msg)
                self.log(f"网关已分发邮件：{subj}")
                break
            except Exception as e:
                if attempt < 2:
                    self.log(f"邮件发送失败，正在重试 ({attempt+1}/3)... 异常: {str(e)}")
                    time.sleep(3)
                else:
                    self.log(f"网络异常导致递交失败：{str(e)}")

    def open_settings_window(self): SettingsWindow(self)

# ===================== 添加/编辑监控项弹窗 =====================
class AddEditWindow(tk.Toplevel):
    def __init__(self, parent, title, initial=None):
        super().__init__(parent)
        self.title(title)
        self.geometry("460x360")
        self.configure(bg="#FFFFFF")
        self.transient(parent)
        self.grab_set()
        self.result, self.init_data = None, initial or {}
        
        frm = tk.Frame(self, bg="#FFFFFF")
        frm.pack(expand=True, fill=tk.BOTH, padx=30, pady=20)
        
        self.enable_var = tk.BooleanVar()
        ttk.Checkbutton(frm, text="分配检测任务", variable=self.enable_var).grid(row=0, column=1, sticky="w", pady=(15, 10))
        
        labels = ["资源名称：", "IP地址：", "协议：", "端口：", "备注："]
        for idx, text in enumerate(labels):
            ttk.Label(frm, text=text, background="#FFFFFF", foreground="#4B5563").grid(row=idx+1, column=0, sticky="e", pady=8, padx=10)
            
        self.name_in = ttk.Entry(frm, width=30)
        self.name_in.grid(row=1, column=1, sticky="w")
        self.addr_in = ttk.Entry(frm, width=30)
        self.addr_in.grid(row=2, column=1, sticky="w")
        
        self.method_var = tk.StringVar()
        self.method_cbx = ttk.Combobox(frm, textvariable=self.method_var, values=["ping", "tcp", "arp"], state="readonly", width=28)
        self.method_cbx.grid(row=3, column=1, sticky="w")
        self.method_cbx.bind("<<ComboboxSelected>>", self.toggle_port)
        
        self.port_in = ttk.Entry(frm, width=30)
        self.port_in.grid(row=4, column=1, sticky="w")
        self.note_in = ttk.Entry(frm, width=30)
        self.note_in.grid(row=5, column=1, sticky="w")
        
        btn_frm = tk.Frame(frm, bg="#FFFFFF")
        btn_frm.grid(row=6, column=0, columnspan=2, pady=25)
        ttk.Button(btn_frm, text="保存修改", command=self.confirm, width=14, style="Primary.TButton").pack(side=tk.LEFT, padx=10)
        ttk.Button(btn_frm, text="取消", command=self.destroy, width=12).pack(side=tk.LEFT)
        
        self.fill_data()
        self.toggle_port()
        self.wait_window()

    def fill_data(self):
        self.enable_var.set(self.init_data.get("enabled", True))
        self.name_in.insert(0, self.init_data.get("name", ""))
        self.addr_in.insert(0, self.init_data.get("address", ""))
        self.method_var.set(self.init_data.get("method", "ping"))
        self.port_in.insert(0, self.init_data.get("port", ""))
        self.note_in.insert(0, self.init_data.get("note", ""))

    def toggle_port(self, evt=None):
        if self.method_var.get() == "tcp": self.port_in.config(state=tk.NORMAL)
        else: self.port_in.delete(0, tk.END); self.port_in.config(state=tk.DISABLED)

    def confirm(self):
        name, addr, method, port, note, en = self.name_in.get().strip(), self.addr_in.get().strip(), self.method_var.get(), self.port_in.get().strip(), self.note_in.get().strip(), self.enable_var.get()
        if not name or not addr: return messagebox.showerror("表单阻断", "标识与地址属性不可为空。")
        if method == "tcp":
            if not port: return messagebox.showerror("表单阻断", "TCP 套接字必须指定目标端口。")
            try:
                if not (1 <= int(port) <= 65535): raise ValueError
            except ValueError: return messagebox.showerror("参数越界", "端口边界许可：1 ~ 65535。")
        self.result = {"name":name, "address":addr, "method":method, "port":port if method=="tcp" else "", "note":note, "enabled":en}
        self.destroy()

# ===================== 设置窗口 =====================
class SettingsWindow(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.parent = parent
        self.title("首选项配置")
        self.geometry("560x600")
        self.minsize(520, 560)
        self.configure(bg="#FFFFFF")
        self.transient(parent)
        self.grab_set()
        self.create_ui()
        self.load_cfg()
        self.wait_window()

    def create_ui(self):
        main_frm = tk.Frame(self, bg="#FFFFFF", padx=20, pady=15)
        main_frm.pack(fill=tk.BOTH, expand=True)
        
        notebook = ttk.Notebook(main_frm)
        notebook.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        
        # ----------------- Tab 1: 性能与调度 -----------------
        mon_tab = tk.Frame(notebook, bg="#FFFFFF", padx=20, pady=20)
        notebook.add(mon_tab, text=" 性能与调度 ")
        
        self.interval_in = ttk.Entry(mon_tab, width=15)
        self.workers_in = ttk.Entry(mon_tab, width=15)
        self.retry_in = ttk.Entry(mon_tab, width=15)
        self.timeout_in = ttk.Entry(mon_tab, width=15)
        self.alert_gap_in = ttk.Entry(mon_tab, width=15)
        
        params = [
            ("轮询间隔 (秒)", self.interval_in),
            ("并发线程数", self.workers_in),
            ("丢包重试阈值", self.retry_in),
            ("探测超时 (秒)", self.timeout_in),
            ("持续离线重告警间隔 (分)", self.alert_gap_in)
        ]
                   
        for idx, (label, entry) in enumerate(params):
            ttk.Label(mon_tab, text=label, background="#FFFFFF", foreground="#5C5C5C").grid(row=idx, column=0, sticky="e", pady=8, padx=10)
            entry.grid(row=idx, column=1, sticky="w")
        
        self.autostart_var = tk.BooleanVar()
        ttk.Checkbutton(mon_tab, text="开机自启 (跟随系统自动后台运行)", variable=self.autostart_var).grid(row=5, column=0, columnspan=2, sticky="w", pady=(20, 0), padx=10)
        
        # ----------------- Tab 2: 告警推送 -----------------
        alert_tab = tk.Frame(notebook, bg="#FFFFFF", padx=20, pady=15)
        notebook.add(alert_tab, text=" 告警与通知 ")
        
        # 桌面通知区域
        self.desktop_notify_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(alert_tab, text="开启系统托盘气泡弹窗通知 (即时桌面提醒)", variable=self.desktop_notify_var).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 10))
        ttk.Separator(alert_tab, orient="horizontal").grid(row=1, column=0, columnspan=3, sticky="ew", pady=(0, 10))

        # Webhook 区域
        self.webhook_en_var = tk.BooleanVar()
        ttk.Checkbutton(alert_tab, text="启用 Webhook 群机器人通知 (企业微信/钉钉/飞书)", variable=self.webhook_en_var).grid(row=2, column=0, columnspan=3, sticky="w", pady=(0, 8))
        
        ttk.Label(alert_tab, text="机器人平台", background="#FFFFFF", foreground="#5C5C5C").grid(row=3, column=0, sticky="e", pady=5, padx=8)
        self.webhook_type_cbx = ttk.Combobox(alert_tab, values=["wecom", "dingtalk", "feishu"], state="readonly", width=15)
        self.webhook_type_cbx.grid(row=3, column=1, sticky="w")
        
        ttk.Label(alert_tab, text="Webhook 地址", background="#FFFFFF", foreground="#5C5C5C").grid(row=4, column=0, sticky="e", pady=5, padx=8)
        self.webhook_url_in = ttk.Entry(alert_tab, width=32)
        self.webhook_url_in.grid(row=4, column=1, columnspan=2, sticky="we")

        ttk.Separator(alert_tab, orient="horizontal").grid(row=5, column=0, columnspan=3, sticky="ew", pady=12)

        # 邮件推送区域
        self.email_en_var = tk.BooleanVar()
        ttk.Checkbutton(alert_tab, text="启用 SMTP 邮件推送", variable=self.email_en_var).grid(row=6, column=0, columnspan=3, sticky="w", pady=(0, 8))
        
        ttk.Label(alert_tab, text="SMTP 服务器", background="#FFFFFF", foreground="#5C5C5C").grid(row=7, column=0, sticky="e", pady=4, padx=8)
        self.smtp_srv = ttk.Entry(alert_tab, width=24)
        self.smtp_srv.grid(row=7, column=1, sticky="w")
        
        port_frm = tk.Frame(alert_tab, bg="#FFFFFF")
        port_frm.grid(row=7, column=2, sticky="w")
        ttk.Label(port_frm, text="端口:").pack(side=tk.LEFT)
        self.smtp_port = ttk.Entry(port_frm, width=6)
        self.smtp_port.pack(side=tk.LEFT, padx=4)
        self.ssl_var = tk.BooleanVar()
        ttk.Checkbutton(port_frm, text="SSL", variable=self.ssl_var).pack(side=tk.LEFT)
        
        ttk.Label(alert_tab, text="发件人邮箱", background="#FFFFFF", foreground="#5C5C5C").grid(row=8, column=0, sticky="e", pady=4, padx=8)
        self.sender_in = ttk.Entry(alert_tab, width=32)
        self.sender_in.grid(row=8, column=1, columnspan=2, sticky="we")
        
        ttk.Label(alert_tab, text="密码/授权码", background="#FFFFFF", foreground="#5C5C5C").grid(row=9, column=0, sticky="e", pady=4, padx=8)
        self.pwd_in = ttk.Entry(alert_tab, show="*", width=32)
        self.pwd_in.grid(row=9, column=1, columnspan=2, sticky="we")
        
        ttk.Label(alert_tab, text="收件人邮箱", background="#FFFFFF", foreground="#5C5C5C").grid(row=10, column=0, sticky="e", pady=4, padx=8)
        self.recv_in = ttk.Entry(alert_tab, width=32)
        self.recv_in.grid(row=10, column=1, columnspan=2, sticky="we")

        ttk.Label(alert_tab, text="抄送邮箱", background="#FFFFFF", foreground="#5C5C5C").grid(row=11, column=0, sticky="e", pady=4, padx=8)
        self.cc_in = ttk.Entry(alert_tab, width=32)
        self.cc_in.grid(row=11, column=1, columnspan=2, sticky="we")
        
        quiet_frm = tk.Frame(alert_tab, bg="#FFFFFF")
        quiet_frm.grid(row=12, column=0, columnspan=3, sticky="w", pady=(8, 0))
        self.email_quiet_var = tk.BooleanVar()
        ttk.Checkbutton(quiet_frm, text="免打扰时段：", variable=self.email_quiet_var).pack(side=tk.LEFT)
        hour_list = [str(i) for i in range(24)]
        self.email_start_cbx = ttk.Combobox(quiet_frm, values=hour_list, width=3, state="readonly")
        self.email_start_cbx.pack(side=tk.LEFT)
        ttk.Label(quiet_frm, text=" 时 ~ ").pack(side=tk.LEFT)
        self.email_end_cbx = ttk.Combobox(quiet_frm, values=hour_list, width=3, state="readonly")
        self.email_end_cbx.pack(side=tk.LEFT)
        ttk.Label(quiet_frm, text=" 时").pack(side=tk.LEFT)
        
        # 底部按钮
        btn_frm = tk.Frame(main_frm, bg="#FFFFFF")
        btn_frm.pack(fill=tk.X, pady=(10, 0))
        ttk.Button(btn_frm, text="保存设置", command=self.save_cfg, width=12, style="Primary.TButton").pack(side=tk.RIGHT, padx=5)
        ttk.Button(btn_frm, text="取消", command=self.destroy, width=10).pack(side=tk.RIGHT)

    def load_cfg(self):
        cfg = self.parent.app_config
        mon = cfg.get("monitoring", {})
        email = cfg.get("email", {})
        sys_cfg = cfg.get("system", {})
        alert = cfg.get("alert", {})
        
        self.interval_in.insert(0, mon.get("interval_seconds", 60))
        self.workers_in.insert(0, mon.get("max_concurrent_workers", 50))
        self.retry_in.insert(0, mon.get("retries_before_offline", 3))
        self.timeout_in.insert(0, mon.get("timeout_seconds", 3))
        self.alert_gap_in.insert(0, mon.get("re_alert_interval_minutes", 60))
        
        self.desktop_notify_var.set(alert.get("desktop_notify", True))
        self.webhook_en_var.set(alert.get("webhook_enabled", False))
        self.webhook_type_cbx.set(alert.get("webhook_type", "wecom"))
        self.webhook_url_in.insert(0, alert.get("webhook_url", ""))
        
        self.email_en_var.set(email.get("enabled", False))
        self.smtp_srv.insert(0, email.get("smtp_server", "smtp.example.com"))
        self.smtp_port.insert(0, email.get("smtp_port", 465))
        self.ssl_var.set(email.get("use_ssl", True))
        self.sender_in.insert(0, email.get("sender_email", ""))
        self.pwd_in.insert(0, email.get("password", ""))
        self.recv_in.insert(0, email.get("recipient_email", ""))
        self.cc_in.insert(0, email.get("cc_email", ""))
        self.email_quiet_var.set(email.get("quiet_hours_enabled", False))
        self.email_start_cbx.set(str(email.get("quiet_start_hour", 22)))
        self.email_end_cbx.set(str(email.get("quiet_end_hour", 8)))
        
        self.autostart_var.set(sys_cfg.get("auto_start", False))

    def save_cfg(self):
        try:
            cfg = self.parent.app_config
            cfg.setdefault("monitoring", {}).update({
                "interval_seconds": int(self.interval_in.get().strip()),
                "max_concurrent_workers": int(self.workers_in.get().strip()),
                "retries_before_offline": int(self.retry_in.get().strip()),
                "timeout_seconds": int(self.timeout_in.get().strip()),
                "re_alert_interval_minutes": int(self.alert_gap_in.get().strip())
            })
            cfg.setdefault("alert", {}).update({
                "desktop_notify": self.desktop_notify_var.get(),
                "webhook_enabled": self.webhook_en_var.get(),
                "webhook_type": self.webhook_type_cbx.get(),
                "webhook_url": self.webhook_url_in.get().strip()
            })
            cfg.setdefault("email", {}).update({
                "enabled": self.email_en_var.get(),
                "smtp_server": self.smtp_srv.get().strip(),
                "smtp_port": int(self.smtp_port.get().strip()),
                "use_ssl": self.ssl_var.get(),
                "sender_email": self.sender_in.get().strip(),
                "password": self.pwd_in.get().strip(),
                "recipient_email": self.recv_in.get().strip(),
                "cc_email": self.cc_in.get().strip(),
                "quiet_hours_enabled": self.email_quiet_var.get(),
                "quiet_start_hour": int(self.email_start_cbx.get()),
                "quiet_end_hour": int(self.email_end_cbx.get())
            })
            auto_start = self.autostart_var.get()
            cfg.setdefault("system", {})["auto_start"] = auto_start
            set_auto_start(auto_start)
            self.parent.save_config()
            self.destroy()
        except ValueError:
            messagebox.showerror("中断", "发现异常输入参数，请核实填写的数据类型。")

# ===================== 程序入口 =====================
if __name__ == '__main__':
    # 启用 Windows DPI 缩放优化，解决高分屏字体界面模糊
    if platform.system() == "Windows":
        import ctypes
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            try:
                ctypes.windll.user32.SetProcessDPIAware()
            except Exception:
                pass

    try: 
        import pystray, PIL, openpyxl
    except ImportError as e:
        messagebox.showerror("运行库阻断", f"缺失模块依赖：{str(e)}\n恢复命令：pip install -r requirements.txt")
        sys.exit(1)
        
    # 单实例锁：优先使用 Windows 内核命名互斥体
    single_instance_handle = acquire_single_instance_lock()
    
    try:
        app = IPMonitorApp()
        app.update_idletasks()
        app.mainloop()
    except KeyboardInterrupt:
        if 'app' in locals() and app.tray_icon: app.tray_icon.stop()
        sys.exit(0)
    except Exception as e:
        messagebox.showerror("异常崩溃", f"运行时发生未捕捉错误：\n{str(e)}")
